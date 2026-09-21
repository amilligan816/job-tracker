"""Turn uploaded files and fetched web pages into plain text for the assistant."""

import io
import logging
import re

import httpx
from fastapi.concurrency import run_in_threadpool
from selectolax.parser import HTMLParser

from app.ats import AtsPosting, resolve_posting

logger = logging.getLogger(__name__)

MAX_FETCH_BYTES = 5 * 1024 * 1024
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_WS = re.compile(r"\n{3,}")


def _collapse(text: str) -> str:
    lines = [line.rstrip() for line in text.splitlines()]
    return _WS.sub("\n\n", "\n".join(lines)).strip()


def html_to_text(html: str) -> str:
    tree = HTMLParser(html)
    for tag in ("script", "style", "noscript", "nav", "footer", "header", "svg"):
        for node in tree.css(tag):
            node.decompose()
    body = tree.body or tree.root
    return _collapse(body.text(separator="\n") if body else "")


async def fetch_posting_text(url: str) -> str:
    """Fetch a job posting URL and reduce it to readable text.

    A career page that is only a shell around an ATS gets resolved through that
    ATS's API first, because scraping the shell returns navigation and legal
    boilerplate rather than the posting. Boards we cannot resolve that way are
    scraped; when those render client-side the text comes back thin and the
    caller should fall back to a pasted description.
    """
    posting = await resolve_posting(url)
    if posting is not None:
        return _render_ats_posting(posting)

    async with httpx.AsyncClient(
        follow_redirects=True,
        timeout=20.0,
        headers={"User-Agent": "job-tracker/0.1 (+personal job search assistant)"},
    ) as client:
        response = await client.get(url)
        response.raise_for_status()
        content = response.content[:MAX_FETCH_BYTES]
        content_type = response.headers.get("content-type", "")

    if "html" in content_type:
        return html_to_text(content.decode(response.encoding or "utf-8", errors="replace"))
    return _collapse(content.decode("utf-8", errors="replace"))


def _render_ats_posting(posting: AtsPosting) -> str:
    """Flatten an ATS payload into the same shape of text a good scrape produces."""
    header = [posting.title]
    for label, value in (
        ("Company", posting.company),
        ("Hiring entity", posting.hiring_entity),
        ("Location", posting.location),
        ("Workplace", posting.workplace),
        ("Employment type", posting.employment_type),
        ("Compensation", posting.compensation),
    ):
        if value:
            header.append(f"{label}: {value}")

    body = html_to_text(posting.content_html)
    return _collapse("\n".join(header) + "\n\n" + body)


async def extract_document_text(data: bytes, content_type: str, filename: str) -> str | None:
    """Best-effort plain text from an upload. Returns None for unsupported formats."""
    lowered = filename.lower()

    if content_type == "application/pdf" or lowered.endswith(".pdf"):
        return await run_in_threadpool(_pdf_to_text, data)

    if content_type.startswith("text/") or lowered.endswith((".txt", ".md", ".markdown")):
        return _collapse(data.decode("utf-8", errors="replace"))

    if "html" in content_type or lowered.endswith((".html", ".htm")):
        return html_to_text(data.decode("utf-8", errors="replace"))

    if lowered.endswith(".docx") or content_type == DOCX_MEDIA_TYPE:
        return await run_in_threadpool(_docx_to_text, data)

    # Legacy .doc is a different, binary format -- store the file, skip the text.
    return None


def _docx_to_text(data: bytes) -> str | None:
    """Plain text from a .docx, including text laid out in tables.

    Walking `document.paragraphs` alone misses table cells, and plenty of
    resumes put their whole layout in an invisible table -- which would look
    like an empty resume to the matcher.
    """
    try:
        from docx import Document as DocxDocument

        document = DocxDocument(io.BytesIO(data))
        blocks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    blocks.extend(p.text for p in cell.paragraphs)

        text = _collapse("\n".join(block for block in blocks if block.strip()))
        return text if len(text) > 40 else None
    except Exception:
        logger.warning("DOCX text extraction failed", exc_info=True)
        return None


def _pdf_to_text(data: bytes) -> str | None:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = _collapse("\n\n".join(pages))
        # A scanned PDF yields near-nothing; treat that as "no text" rather than
        # feeding the model a page of whitespace.
        return text if len(text) > 40 else None
    except Exception:
        logger.warning("PDF text extraction failed", exc_info=True)
        return None
